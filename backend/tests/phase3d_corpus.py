"""Phase 3D evaluation corpus.

Real documents, built programmatically (the same approach as
``tests/test_structured_extraction.py``), covering every document shape the
Phase 3D contract names:

  short sections · long paragraphs · unpunctuated text · small and large
  tables · lists · multi-page sections · heading-heavy documents ·
  academic / two-column PDFs · mixed documents

Each document carries labelled questions.  Labelling rules:

* ``evidence`` entries are short, distinctive substrings that must appear in
  a chunk for it to count as supporting evidence.  Invented proper nouns
  (Zephyr, Quokka, Kestrel, …) are used so a match cannot be accidental.
* ``expected_answer`` is the exact span an extractive model should return
  (short noun phrases only — the QA model is extractive).
* ``answerable=False`` questions have no supporting text anywhere in the
  corpus; they measure abstention, never score.

Instants of truth live in this module only; nothing here depends on the
chunker, so the same corpus is used for both strategies.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import docx
import fitz  # PyMuPDF


# ---------------------------------------------------------------------------
# Types
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Question:
    qid: str
    text: str
    answerable: bool = True
    expected_answer: str | None = None
    evidence: tuple[str, ...] = ()
    kind: str = "factual"
    notes: str = ""


@dataclass
class CorpusDoc:
    name: str
    file_type: str
    builder: Callable[[Path], None]
    questions: list[Question] = field(default_factory=list)
    notes: str = ""

    def write(self, directory: Path) -> Path:
        path = directory / self.name
        self.builder(path)
        return path


# ---------------------------------------------------------------------------
# Builders — DOCX
# ---------------------------------------------------------------------------

def _write_docx(path: Path, add) -> None:
    document = docx.Document()
    add(document)
    document.save(str(path))


def _build_short_sections(path: Path) -> None:
    def add(document):
        document.add_heading("Reference Manual", level=1)
        sections = [
            ("Alpha", "The Zephyr valve is calibrated at 22 degrees."),
            ("Beta", "The Nimbus sensor reports in millibars."),
            ("Gamma", "The Helios lamp draws 14 watts."),
            ("Delta", "The Quokka bracket weighs 3 kilograms."),
            ("Epsilon", "The Vermilion coating cures in 6 hours."),
            ("Zeta", "The Solstice bearing tolerates 300 degrees."),
        ]
        for title, body in sections:
            document.add_heading(title, level=2)
            document.add_paragraph(body)

    _write_docx(path, add)


def _build_heading_heavy(path: Path) -> None:
    def add(document):
        document.add_heading("Handbook", level=1)
        document.add_heading("Installation", level=2)
        document.add_heading("Preparation", level=3)
        document.add_heading("Tools", level=4)
        document.add_paragraph("A Tundra torque wrench is required.")
        document.add_heading("Wiring", level=4)
        document.add_paragraph("Route the Kestrel cable behind the panel.")
        document.add_heading("Commissioning", level=2)
        document.add_paragraph("Commission the unit within 30 days.")
        document.add_heading("Maintenance", level=1)
        document.add_paragraph("Replace the Zephyr filter every 500 hours.")

    _write_docx(path, add)


def _build_tables_small(path: Path) -> None:
    def add(document):
        document.add_heading("Results", level=1)
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Method"
        table.cell(0, 1).text = "Result"
        table.cell(1, 0).text = "PCA"
        table.cell(1, 1).text = "85 percent"
        table.cell(2, 0).text = "SVM"
        table.cell(2, 1).text = "92 percent"

    _write_docx(path, add)


def _build_tables_large(path: Path) -> None:
    def add(document):
        document.add_heading("Results", level=1)
        table = document.add_table(rows=31, cols=3)
        table.cell(0, 0).text = "Method"
        table.cell(0, 1).text = "Accuracy"
        table.cell(0, 2).text = "Runtime"
        for i in range(1, 31):
            table.cell(i, 0).text = f"Method {i}"
            table.cell(i, 1).text = f"{60 + i} percent"
            table.cell(i, 2).text = f"{i * 7} milliseconds"
        # Method 30 is the row that only a header-carrying sub-chunk can explain.
        add_para = document.add_paragraph
        add_para("The recommended configuration is Method 30.")

    _write_docx(path, add)


def _build_lists(path: Path) -> None:
    def add(document):
        document.add_heading("Procedure", level=1)
        for item in [
            "Isolate the Nimbus supply",
            "Drain the Vermilion loop",
            "Remove the Helios filter",
            "Torque the Quokka bolts",
        ]:
            document.add_paragraph(item, style="List Bullet")
        document.add_heading("Checks", level=2)
        for item in [
            "Confirm the Zephyr reading",
            "Log the Solstice temperature",
        ]:
            document.add_paragraph(item, style="List Number")

    _write_docx(path, add)


def _build_mixed(path: Path) -> None:
    def add(document):
        document.add_heading("Observatory Report", level=1)
        document.add_paragraph(
            "The Tundra survey recorded atmospheric pressure across three "
            "seasons using the Nimbus array."
        )
        document.add_heading("Instrumentation", level=2)
        document.add_paragraph(
            "The Nimbus array uses a Quokka reference cell for calibration."
        )
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Season"
        table.cell(0, 1).text = "Pressure"
        table.cell(1, 0).text = "Spring"
        table.cell(1, 1).text = "1013 millibars"
        table.cell(2, 0).text = "Winter"
        table.cell(2, 1).text = "998 millibars"
        document.add_heading("Findings", level=2)
        document.add_paragraph(
            "The lowest recorded pressure was 998 millibars."
        )
        for item in [
            "The Helios sensor drifted by 2 millibars",
            "The Zephyr probe remained stable",
        ]:
            document.add_paragraph(item, style="List Bullet")

    _write_docx(path, add)


# ---------------------------------------------------------------------------
# Builders — TXT
# ---------------------------------------------------------------------------

def _build_long_paragraph(path: Path) -> None:
    sentences = [
        "The Meridian programme was established to study coastal erosion.",
        "Its first survey was completed by the Kestrel team in 1988.",
        "The team measured sediment transport along twelve beaches.",
        "Each beach was sampled at three tidal heights.",
        "The samples were analysed using laser diffraction.",
        "The resulting dataset contained over forty thousand records.",
        "A follow-up survey was commissioned in 1994.",
        "The second survey extended coverage to twenty beaches.",
        "The Kestrel team reported a mean erosion rate.",
        "The reported mean erosion rate was 1.4 metres per year.",
        "The findings were published in the Meridian bulletin.",
        "The bulletin recommended annual monitoring.",
    ]
    paragraphs = [" ".join(sentences)] * 2
    path.write_text("\n\n".join(paragraphs) + "\n", encoding="utf-8")


def _build_dense_prose(path: Path) -> None:
    """A single long, prose-heavy report.

    This is the shape that exposes the context-window defect: it is large
    enough that the legacy chunker emits many ~1000-character chunks, so a
    top-5 retrieval supplies far more text than the QA model can read.
    """
    paragraphs = [
        (
            "The Mariner survey was commissioned to characterise sediment "
            "transport along the eastern shelf. Fieldwork began in 2009 and "
            "continued for four consecutive seasons. Each season covered a "
            "different transect, with sampling repeated at three tidal states. "
            "The survey used a Kestrel frame to position every sampling station."
        ),
        (
            "Sampling followed a stratified random design. Stations were "
            "allocated in proportion to the estimated bed area of each stratum. "
            "This allocation reduced the variance of the regional sediment "
            "estimates compared with a simple random design. The design was "
            "reviewed by an independent statistical panel before fieldwork began."
        ),
        (
            "Grain size was determined by laser diffraction. Each sample was "
            "split into three subsamples and measured separately. The median "
            "grain size was recorded for each subsample and averaged afterwards. "
            "Repeat measurements agreed to within two percent of the median. "
            "The diffraction instrument was calibrated weekly against a standard."
        ),
        (
            "The survey recorded a total of 41200 measurements. Transport rates "
            "were derived from the grain-size distributions using a calibrated "
            "bedload function. The derived rates carried an uncertainty of "
            "roughly fifteen percent. That uncertainty was dominated by the "
            "choice of bedload function rather than by measurement error."
        ),
        (
            "The mean transport rate was 1.4 kilograms per metre per day. Rates "
            "varied strongly between transects, with the northern transect "
            "carrying almost twice the regional mean. The southern transect "
            "carried the least material of any transect in the survey. "
            "These differences persisted across all four seasons."
        ),
        (
            "A second objective was to assess the stability of the Kestrel "
            "frame under load. Frame deflection was logged continuously during "
            "deployment. The maximum recorded deflection was 14 millimetres. "
            "Deflection of that magnitude was observed only twice, both in "
            "storm conditions."
        ),
        (
            "Sediment cores were collected at twelve reference stations. Each "
            "core was logged for stratigraphy before subsampling. The cores "
            "revealed a consistent fining-upward sequence. Radiometric dating "
            "placed the base of the sequence at approximately 1850."
        ),
        (
            "The survey concluded with a set of management recommendations. "
            "The panel recommended annual monitoring of the northern transect. "
            "It further recommended that the Kestrel frame be replaced before "
            "the next survey cycle. The final recommendation concerned the "
            "archiving of the raw diffraction measurements."
        ),
    ]
    path.write_text("\n\n".join(paragraphs) + "\n", encoding="utf-8")


def _build_unpunctuated(path: Path) -> None:
    # Deliberately unpunctuated: exercises the hard-cut path.
    path.write_text("AlphaBetaGammaDelta " * 120 + "\n", encoding="utf-8")


def _build_markdown_sections(path: Path) -> None:
    path.write_text(
        "# Overview\n\n"
        "The Tundra station opened in 2005.\n\n"
        "## Power\n\n"
        "The station uses a Helios generator rated at 40 kilowatts.\n\n"
        "## Water\n\n"
        "Water is drawn from the Vermilion aquifer.\n\n"
        "# Decommissioning\n\n"
        "The station closed in 2019.\n",
        encoding="utf-8",
    )


# ---------------------------------------------------------------------------
# Builders — PDF
# ---------------------------------------------------------------------------

def _pdf_text(page, x: float, y: float, text: str, size: float = 11.0, bold: bool = False):
    font = fitz.Font("hebo" if bold else "helv")
    writer = fitz.TextWriter(page.rect)
    writer.append((x, y), text, font=font, fontsize=size)
    writer.write_text(page)


def _build_multipage_section(path: Path) -> None:
    document = fitz.open()
    page1 = document.new_page(width=595, height=842)
    _pdf_text(page1, 50, 60, "Orbital Mechanics", size=16, bold=True)
    _pdf_text(page1, 50, 90, "The Kestrel programme began in 1971.")
    _pdf_text(page1, 50, 110, "Its first launch used the Nimbus booster.")

    page2 = document.new_page(width=595, height=842)
    _pdf_text(page2, 50, 60, "The programme entered its second phase.")
    _pdf_text(page2, 50, 80, "The Falcon array replaced the Nimbus booster.")
    _pdf_text(page2, 50, 100, "Three flights were conducted that decade.")

    page3 = document.new_page(width=595, height=842)
    _pdf_text(page3, 50, 60, "The Kestrel programme concluded in 1984.")
    _pdf_text(page3, 50, 80, "The Falcon array was retired after the final flight.")
    document.save(str(path))
    document.close()


def _build_academic_two_column(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    # Left column
    _pdf_text(page, 50, 60, "Abstract", size=14, bold=True)
    _pdf_text(page, 50, 85, "We present the Quokka estimator for sparse")
    _pdf_text(page, 50, 102, "spectral recovery. The estimator builds on")
    _pdf_text(page, 50, 119, "the Tundra basis and converges quickly.")
    # Right column (must not interleave with the left column)
    _pdf_text(page, 320, 60, "1. Introduction", size=14, bold=True)
    _pdf_text(page, 320, 85, "Prior work used the Vermilion transform")
    _pdf_text(page, 320, 102, "with limited success on noisy inputs.")
    _pdf_text(page, 320, 119, "Our contribution improves recall by 11 percent.")
    document.save(str(path))
    document.close()


# ---------------------------------------------------------------------------
# Corpus
# ---------------------------------------------------------------------------

def build_corpus() -> list[CorpusDoc]:
    """Return the Phase 3D evaluation corpus (documents + labelled questions)."""
    return [
        CorpusDoc(
            name="short_sections.docx",
            file_type="docx",
            builder=_build_short_sections,
            notes="Six short H2 sections — target for short-element merging.",
            questions=[
                Question("ss_1", "At what angle is the Zephyr valve calibrated?",
                         expected_answer="22 degrees", evidence=("Zephyr",), kind="factual"),
                Question("ss_2", "How many watts does the Helios lamp draw?",
                         expected_answer="14 watts", evidence=("Helios",), kind="factual"),
                Question("ss_3", "How much does the Quokka bracket weigh?",
                         expected_answer="3 kilograms", evidence=("Quokka",), kind="factual"),
                Question("ss_4", "What does the Epsilon section describe?",
                         expected_answer="Vermilion", evidence=("Vermilion", "6 hours"),
                         kind="section_specific"),
            ],
        ),
        CorpusDoc(
            name="heading_heavy.docx",
            file_type="docx",
            builder=_build_heading_heavy,
            notes="H1–H4 nesting — heading prefix cost and section scoping.",
            questions=[
                Question("hh_1", "What tool is required for preparation?",
                         expected_answer="Tundra torque wrench", evidence=("Tundra",),
                         kind="section_specific"),
                Question("hh_2", "Where should the Kestrel cable be routed?",
                         expected_answer="behind the panel", evidence=("Kestrel",),
                         kind="section_specific"),
                Question("hh_3", "How often should the Zephyr filter be replaced?",
                         expected_answer="every 500 hours", evidence=("Zephyr",),
                         kind="factual"),
            ],
        ),
        CorpusDoc(
            name="tables_small.docx",
            file_type="docx",
            builder=_build_tables_small,
            notes="Small table — must stay atomic.",
            questions=[
                Question("ts_1", "What result did SVM achieve?",
                         expected_answer="92 percent", evidence=("SVM", "92 percent"),
                         kind="table"),
                Question("ts_2", "What result did PCA achieve?",
                         expected_answer="85 percent", evidence=("PCA", "85 percent"),
                         kind="table"),
            ],
        ),
        CorpusDoc(
            name="tables_large.docx",
            file_type="docx",
            builder=_build_tables_large,
            notes="31-row table — row split with the header repeated.",
            questions=[
                Question("tl_1", "What is the accuracy of Method 30?",
                         expected_answer="90 percent", evidence=("Method 30",),
                         kind="table"),
                Question("tl_2", "Which configuration is recommended?",
                         expected_answer="Method 30", evidence=("Method 30",),
                         kind="table"),
                Question("tl_3", "What runtime does Method 12 have?",
                         expected_answer="84 milliseconds", evidence=("Method 12",),
                         kind="table"),
            ],
        ),
        CorpusDoc(
            name="lists.docx",
            file_type="docx",
            builder=_build_lists,
            notes="Bulleted then numbered list — item integrity and ordering.",
            questions=[
                Question("li_1", "Which item follows draining the Vermilion loop?",
                         expected_answer="Remove the Helios filter",
                         evidence=("Vermilion", "Helios"), kind="list"),
                Question("li_2", "What should be isolated first?",
                         expected_answer="Nimbus supply", evidence=("Nimbus",), kind="list"),
                Question("li_3", "What must be confirmed in the checks section?",
                         expected_answer="Zephyr reading", evidence=("Zephyr",), kind="list"),
            ],
        ),
        CorpusDoc(
            name="mixed.docx",
            file_type="docx",
            builder=_build_mixed,
            notes="Prose + heading + table + list in one document.",
            questions=[
                Question("mx_1", "What was the lowest recorded pressure?",
                         expected_answer="998 millibars",
                         evidence=("998 millibars",), kind="table"),
                Question("mx_2", "What does the Nimbus array use for calibration?",
                         expected_answer="Quokka reference cell",
                         evidence=("Quokka",), kind="section_specific"),
                Question("mx_3", "Which season recorded 1013 millibars?",
                         expected_answer="Spring", evidence=("1013 millibars",),
                         kind="table"),
                Question("mx_4", "How far did the Helios sensor drift?",
                         expected_answer="2 millibars", evidence=("Helios",),
                         kind="list"),
            ],
        ),
        CorpusDoc(
            name="long_paragraph.txt",
            file_type="txt",
            builder=_build_long_paragraph,
            notes="Two very long paragraphs — sentence splitting and overlap.",
            questions=[
                Question("lp_1", "What was the reported mean erosion rate?",
                         expected_answer="1.4 metres per year",
                         evidence=("1.4 metres per year",), kind="long_paragraph"),
                Question("lp_2", "In which year was the first survey completed?",
                         expected_answer="1988", evidence=("Kestrel",), kind="factual"),
                Question("lp_3", "How many beaches did the second survey cover?",
                         expected_answer="twenty beaches",
                         evidence=("twenty beaches",), kind="long_paragraph"),
            ],
        ),
        CorpusDoc(
            name="dense_prose.txt",
            file_type="txt",
            builder=_build_dense_prose,
            notes="Prose-heavy report — the shape that overflows the QA window.",
            questions=[
                Question("dp_1", "What instrument determined the grain size?",
                         expected_answer="laser diffraction",
                         evidence=("laser diffraction",), kind="dense_prose"),
                Question("dp_2", "What was the mean transport rate?",
                         expected_answer="1.4 kilograms per metre per day",
                         evidence=("1.4 kilograms per metre per day",),
                         kind="dense_prose"),
                Question("dp_3", "What was the maximum recorded frame deflection?",
                         expected_answer="14 millimetres",
                         evidence=("14 millimetres",), kind="dense_prose"),
                Question("dp_4", "When was the base of the sequence dated to?",
                         expected_answer="1850", evidence=("1850",),
                         kind="dense_prose"),
                Question("dp_5", "Which transect should be monitored annually?",
                         expected_answer="northern transect",
                         evidence=("northern transect",), kind="dense_prose"),
                Question("dp_6", "What was reviewed before fieldwork began?",
                         expected_answer="design",
                         evidence=("independent statistical panel",),
                         kind="surrounding_context"),
            ],
        ),
        CorpusDoc(
            name="unpunctuated.txt",
            file_type="txt",
            builder=_build_unpunctuated,
            notes="Unpunctuated text — exercises the hard-cut path honestly.",
            questions=[
                Question("up_1", "Which token repeats in the document?",
                         expected_answer="AlphaBetaGammaDelta",
                         evidence=("AlphaBetaGammaDelta",), kind="unpunctuated"),
            ],
        ),
        CorpusDoc(
            name="markdown_sections.txt",
            file_type="txt",
            builder=_build_markdown_sections,
            notes="Markdown headings in TXT — heading context from # levels.",
            questions=[
                Question("md_1", "When did the Tundra station open?",
                         expected_answer="2005", evidence=("Tundra",),
                         kind="factual"),
                Question("md_2", "What generator does the station use?",
                         expected_answer="Helios generator", evidence=("Helios",),
                         kind="section_specific"),
                Question("md_3", "When did the station close?",
                         expected_answer="2019", evidence=("Decommissioning",),
                         kind="section_specific"),
            ],
        ),
        CorpusDoc(
            name="multipage_section.pdf",
            file_type="pdf",
            builder=_build_multipage_section,
            notes="One section spanning three pages — page boundary behaviour.",
            questions=[
                Question("mp_1", "When did the Kestrel programme begin?",
                         expected_answer="1971", evidence=("1971",), kind="factual"),
                Question("mp_2", "When did the Kestrel programme conclude?",
                         expected_answer="1984", evidence=("1984",), kind="multipage"),
                Question("mp_3", "What replaced the Nimbus booster?",
                         expected_answer="Falcon array", evidence=("Falcon",),
                         kind="multipage"),
                Question("mp_4", "Trace the Kestrel programme from start to finish.",
                         expected_answer="1984", evidence=("1971", "1984"),
                         kind="multichunk"),
            ],
        ),
        CorpusDoc(
            name="academic_two_column.pdf",
            file_type="pdf",
            builder=_build_academic_two_column,
            notes="Two-column academic PDF — block order and column integrity.",
            questions=[
                Question("ac_1", "What is the name of the estimator?",
                         expected_answer="Quokka estimator", evidence=("Quokka",),
                         kind="academic"),
                Question("ac_2", "Which basis does the estimator build on?",
                         expected_answer="Tundra basis", evidence=("Tundra",),
                         kind="academic"),
                Question("ac_3", "By how much does the contribution improve recall?",
                         expected_answer="11 percent", evidence=("11 percent",),
                         kind="academic"),
            ],
        ),
    ]


def unanswerable_questions() -> list[Question]:
    """Questions with no supporting evidence anywhere in the corpus."""
    return [
        Question("un_1", "What is the capital of Finland?", answerable=False,
                 kind="unsupported"),
        Question("un_2", "How many employees work at the observatory?",
                 answerable=False, kind="unsupported"),
        Question("un_3", "What is the melting point of tungsten?",
                 answerable=False, kind="unsupported"),
    ]
